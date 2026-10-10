"""Render the v2 method framework figure into native, editable PowerPoint shapes.

A minimal matplotlib renderer emits every path as a freeform shape, every text
line as a text box (math converted to sub/superscript runs), and every raster as
a picture, so the result can be edited element by element in PowerPoint.
"""

from __future__ import annotations

import io
import re
import subprocess
import sys
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import numpy as np
from matplotlib.backend_bases import RendererBase
from matplotlib.path import Path as MPath
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml import parse_xml
from pptx.util import Emu, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))
import plot_fig_method_framework as v2  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "figures" / "fig_method_framework_v2_editable.pptx"
EMU_PER_IN = 914400
NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" ' \
     'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'

SYMBOLS = {
    "times": "\u00D7", "in": "\u2208", "rightarrow": "\u2192", "to": "\u2192", "geq": "\u2265",
    "leq": "\u2264", "ell": "\u2113", "pi": "\u03C0", "tau": "\u03C4", "cdot": "\u00B7",
    "|": "\u2016", "Vert": "\u2016", ",": "\u2009", " ": " ", ";": " ", "quad": "  ",
}
SPACED = {"=", "<", ">", "+", "\u2208", "\u2265", "\u2264", "\u2192", "\u00D7"}
ACCENTS = {"hat": "\u0302", "tilde": "\u0303", "bar": "\u0304"}
BLACKBOARD = {"R": "\u211D", "N": "\u2115", "Z": "\u2124"}


def _tokens(s):
    return [t for t in re.findall(r"\\[A-Za-z]+|\\.|[{}^_]|.", s) if not t.isspace()]


def _parse(toks, i, style, out, stop_at_close=False):
    """Append (text, style) runs; style = dict(italic, upright, baseline)."""
    while i < len(toks):
        t = toks[i]
        if t == "}":
            return i + 1
        if t == "{":
            i = _parse(toks, i + 1, style, out, True)
            continue
        if t in ("^", "_"):
            sub = dict(style, baseline=30000 if t == "^" else -25000)
            i += 1
            if i < len(toks) and toks[i] == "{":
                i = _parse(toks, i + 1, sub, out, True)
            else:
                i = _parse(toks[i:i + 1], 0, sub, out) + i
            continue
        if t.startswith("\\"):
            name = t[1:]
            i += 1
            if name in ("mathrm", "text", "mathbf"):
                i = _parse(toks, i + 1, dict(style, upright=True, bold=name == "mathbf"), out, True)
            elif name == "mathbb":
                inner = []
                i = _parse(toks, i + 1, dict(style, upright=True), inner, True)
                out.extend((("".join(BLACKBOARD.get(ch, ch) for ch in txt)), st) for txt, st in inner)
            elif name in ACCENTS:
                inner = []
                if i < len(toks) and toks[i] == "{":
                    i = _parse(toks, i + 1, style, inner, True)
                else:
                    i = _parse(toks[i:i + 1], 0, style, inner) + i
                if inner:
                    txt, st = inner[-1]
                    inner[-1] = (txt + ACCENTS[name], st)
                out.extend(inner)
            else:
                sym = SYMBOLS.get(name, name)
                if sym in SPACED and not stop_at_close and not style.get("baseline"):
                    sym = f" {sym} "
                out.append((sym, dict(style, upright=True)))
            continue
        if t in SPACED and not stop_at_close and not style.get("baseline"):
            t = f" {t} "
        out.append((t, style))
        i += 1
    return i


def math_runs(s):
    """Split a matplotlib string into (text, italic, bold, baseline) runs."""
    runs = []
    for k, part in enumerate(s.split("$")):
        if not part:
            continue
        if k % 2 == 0:
            runs.append((part, None, None, 0))
            continue
        out = []
        _parse(_tokens(part), 0, {"italic": True, "upright": False, "baseline": 0}, out)
        for txt, st in out:
            italic = st.get("italic") and not st.get("upright") and any(c.isalpha() for c in txt)
            runs.append((txt, bool(italic), st.get("bold") or None, st.get("baseline", 0)))
    merged = []
    for r in runs:
        if merged and merged[-1][1:] == r[1:]:
            merged[-1] = (merged[-1][0] + r[0],) + r[1:]
        else:
            merged.append(r)
    return merged


def _hex(rgb):
    return "%02X%02X%02X" % tuple(int(round(max(0, min(1, c)) * 255)) for c in rgb[:3])


def _fill_xml(rgb, alpha):
    a = "" if alpha >= 0.999 else f'<a:alpha val="{int(alpha * 100000)}"/>'
    return f'<a:solidFill><a:srgbClr val="{_hex(rgb)}">{a}</a:srgbClr></a:solidFill>'


class PptxRenderer(RendererBase):
    def __init__(self, slide, width_px, height_px, dpi):
        super().__init__()
        self.slide = slide
        self.width, self.height, self.dpi = width_px, height_px, dpi
        self.next_id = 1000
        self.skip_first_bg = True

    # ---- coordinate helpers
    def emu(self, x, y):
        return int(x / self.dpi * EMU_PER_IN), int((self.height - y) / self.dpi * EMU_PER_IN)

    def flipy(self):
        return False

    def get_canvas_width_height(self):
        return self.width, self.height

    def points_to_pixels(self, points):
        return points * self.dpi / 72.0

    def option_image_nocomposite(self):
        return True

    def option_scale_image(self):
        return False

    # ---- paths
    def draw_path(self, gc, path, transform, rgbFace=None):
        segs = list(path.iter_segments(transform, remove_nans=True, simplify=False, curves=True))
        if not segs:
            return
        pts = np.concatenate([np.asarray(v).reshape(-1, 2) for v, _ in segs])
        x0, y0 = pts[:, 0].min(), pts[:, 1].min()
        x1, y1 = pts[:, 0].max(), pts[:, 1].max()

        lw = gc.get_linewidth()
        stroke_rgba = gc.get_rgb()
        stroke_alpha = gc.get_alpha() if gc.get_forced_alpha() else stroke_rgba[3]
        has_fill = rgbFace is not None and (len(rgbFace) < 4 or rgbFace[3] > 0)
        has_line = lw > 0 and stroke_alpha > 0

        if self.skip_first_bg:
            self.skip_first_bg = False
            if has_fill and x1 - x0 >= self.width - 1 and y1 - y0 >= self.height - 1:
                return
        if not has_fill and not has_line:
            return

        ox, oy = self.emu(x0, y1)
        cx = max(int((x1 - x0) / self.dpi * EMU_PER_IN), 1)
        cy = max(int((y1 - y0) / self.dpi * EMU_PER_IN), 1)

        def p(x, y):
            px, py = self.emu(x, y)
            return f'<a:pt x="{px - ox}" y="{py - oy}"/>'

        cmds = []
        for verts, code in segs:
            v = np.asarray(verts).reshape(-1, 2)
            if code == MPath.MOVETO:
                cmds.append(f"<a:moveTo>{p(*v[0])}</a:moveTo>")
            elif code == MPath.LINETO:
                cmds.append(f"<a:lnTo>{p(*v[0])}</a:lnTo>")
            elif code == MPath.CURVE3:
                cmds.append(f"<a:quadBezTo>{p(*v[0])}{p(*v[1])}</a:quadBezTo>")
            elif code == MPath.CURVE4:
                cmds.append(f"<a:cubicBezTo>{p(*v[0])}{p(*v[1])}{p(*v[2])}</a:cubicBezTo>")
            elif code == MPath.CLOSEPOLY:
                cmds.append("<a:close/>")
        path_fill = "" if has_fill else ' fill="none"'

        if has_fill:
            fa = rgbFace[3] if len(rgbFace) > 3 else 1.0
            if gc.get_forced_alpha():
                fa = gc.get_alpha()
            fill = _fill_xml(rgbFace, fa)
        else:
            fill = "<a:noFill/>"

        if has_line:
            dash = ""
            offset, seq = gc.get_dashes()
            if seq is not None and len(seq) >= 2:
                lw_px = max(lw, 0.1)
                ds = "".join(
                    f'<a:ds d="{int(seq[k] / lw_px * 100000)}" sp="{int(seq[k + 1] / lw_px * 100000)}"/>'
                    for k in range(0, len(seq) - 1, 2))
                dash = f"<a:custDash>{ds}</a:custDash>"
            cap = {"butt": "flat", "round": "rnd", "projecting": "sq"}.get(gc.get_capstyle(), "flat")
            join = {"round": "<a:round/>", "bevel": "<a:bevel/>"}.get(gc.get_joinstyle(), "<a:miter/>")
            line = (f'<a:ln w="{int(lw * 12700)}" cap="{cap}">{_fill_xml(stroke_rgba, stroke_alpha)}'
                    f"{dash}{join}</a:ln>")
        else:
            line = "<a:ln><a:noFill/></a:ln>"

        self.next_id += 1
        xml = (
            f'<p:sp {NS}><p:nvSpPr><p:cNvPr id="{self.next_id}" name="Shape {self.next_id}"/>'
            f"<p:cNvSpPr/><p:nvPr/></p:nvSpPr><p:spPr>"
            f'<a:xfrm><a:off x="{ox}" y="{oy}"/><a:ext cx="{cx}" cy="{cy}"/></a:xfrm>'
            f'<a:custGeom><a:avLst/><a:gdLst/><a:ahLst/><a:cxnLst/><a:rect l="0" t="0" r="r" b="b"/>'
            f'<a:pathLst><a:path w="{cx}" h="{cy}"{path_fill}>{"".join(cmds)}</a:path></a:pathLst>'
            f"</a:custGeom>{fill}{line}</p:spPr></p:sp>"
        )
        self.slide.shapes._spTree.append(parse_xml(xml))

    # ---- images
    def draw_image(self, gc, x, y, im, transform=None):
        h, w = im.shape[:2]
        buf = io.BytesIO()
        Image.fromarray(np.asarray(im)[::-1]).save(buf, format="PNG")
        buf.seek(0)
        left, top = self.emu(x, y + h)
        self.slide.shapes.add_picture(buf, Emu(left), Emu(top),
                                      Emu(int(w / self.dpi * EMU_PER_IN)), Emu(int(h / self.dpi * EMU_PER_IN)))

    # ---- text
    def draw_text(self, gc, x, y, s, prop, angle, ismath=False, mtext=None):
        if not s.strip():
            return
        size = prop.get_size_in_points()
        w_px, h_px, d_px = self.get_text_width_height_descent(s, prop, ismath)
        w_in = w_px / self.dpi * 1.12 + 0.02
        ascent_in = size * 0.905 / 72
        line_in = size * 1.2 / 72
        bx, by = x / self.dpi, (self.height - y) / self.dpi
        left, top = bx, by - ascent_in
        if angle:
            a = np.deg2rad(angle)
            cx_u, cy_u = left + w_in / 2 - bx, top + line_in / 2 - by
            rx = cx_u * np.cos(a) + cy_u * np.sin(a)
            ry = -cx_u * np.sin(a) + cy_u * np.cos(a)
            left, top = bx + rx - w_in / 2, by + ry - line_in / 2

        tb = self.slide.shapes.add_textbox(Emu(int(left * EMU_PER_IN)), Emu(int(top * EMU_PER_IN)),
                                           Emu(int(w_in * EMU_PER_IN)), Emu(int(line_in * EMU_PER_IN)))
        if angle:
            tb.rotation = -angle
        tf = tb.text_frame
        tf.word_wrap = False
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.TOP
        para = tf.paragraphs[0]
        para.alignment = PP_ALIGN.LEFT

        weight = prop.get_weight()
        base_bold = weight == "bold" or (isinstance(weight, (int, float)) and weight >= 600) or weight in (
            "semibold", "demibold", "heavy", "extra bold", "black")
        base_italic = prop.get_style() in ("italic", "oblique")
        color = RGBColor.from_string(_hex(gc.get_rgb()))
        runs = math_runs(s) if ismath else [(s, None, None, 0)]
        for txt, italic, bold, baseline in runs:
            r = para.add_run()
            r.text = txt
            f = r.font
            f.name = "Arial"
            f.size = Pt(size)
            f.bold = base_bold if bold is None else (bold or base_bold)
            f.italic = base_italic if italic is None else italic
            f.color.rgb = color
            if baseline:
                r._r.get_or_add_rPr().set("baseline", str(baseline))

    def draw_tex(self, gc, x, y, s, prop, angle, *, mtext=None):
        self.draw_text(gc, x, y, s, prop, angle, True, mtext)


def main() -> None:
    matplotlib.rcParams["image.composite_image"] = False
    fig = v2.build_figure()
    dpi = 300
    fig.set_dpi(dpi)
    w_in, h_in = fig.get_size_inches()

    prs = Presentation()
    prs.slide_width = Emu(int(w_in * EMU_PER_IN))
    prs.slide_height = Emu(int(h_in * EMU_PER_IN))
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    renderer = PptxRenderer(slide, w_in * dpi, h_in * dpi, dpi)
    fig.draw(renderer)
    prs.save(OUT)
    print(f"wrote {OUT} ({len(slide.shapes)} shapes)")

    png = OUT.with_name(OUT.stem + "_pptx_preview.png")
    ps = (
        "$pp = New-Object -ComObject PowerPoint.Application; "
        f"$p = $pp.Presentations.Open('{OUT}', $true, $false, $false); "
        f"$p.Slides.Item(1).Export('{png}', 'PNG', 2400, {int(2400 * h_in / w_in)}); "
        "$p.Close(); $pp.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    print(f"preview {png}")


if __name__ == "__main__":
    main()
